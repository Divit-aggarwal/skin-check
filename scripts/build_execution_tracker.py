"""PROJECT SKIN execution tracker: single data source -> Markdown + XLSX.

Edit the data below, then regenerate both files:
    uv run --with openpyxl python scripts/build_execution_tracker.py

The script refuses to write if a tracker rule is violated (unknown status, DONE
without evidence, BLOCKED without blocker, missing next action, unknown dependency).
"""
from pathlib import Path

UPDATED = "2026-10-03"
VALIDATED_COMMIT = "91d20bc"
ROOT = Path(__file__).resolve().parent.parent
MD_OUT = ROOT / "docs/PROJECT_SKIN_EXECUTION_TRACKER.md"
XLSX_OUT = ROOT / "docs/PROJECT_SKIN_EXECUTION_TRACKER.xlsx"

STATUSES = ["DONE", "IN PROGRESS", "PREPARED", "UNBLOCKED", "BLOCKED", "RESEARCH",
            "LATER", "OWNER ACTION", "NEEDS VERIFICATION"]
MILESTONES = ["FOUNDATION", "TRUSTED MEASUREMENT", "CALIBRATION", "LONGITUDINAL INTELLIGENCE",
              "ROUTINE/PRODUCT SYSTEM", "GEMINI CONTEXT LAYER", "DASHBOARD", "PRIVACY/SAFETY",
              "RED TEAM", "RELEASE"]

NOT_DEPLOYED = "Implemented on branch longitudinal-platform only. NOT deployed: production frontend is 1b39f88 (master), which contains none of gates 01-04."

# ---------------------------------------------------------------- overview facts
OVERVIEW = [
    ("Project", "PROJECT SKIN / Longitudinal Skin Management Platform"),
    ("Current branch", "longitudinal-platform (worktree also holds infra/india-region at ff08107, unmerged)"),
    ("Current commit", "91d20bc4025d1ad34b1b7de7bcffd9fc488f776b (working tree has uncommitted DECISION_LOG.md edits and untracked DATA_LOCALITY_DECISION.md, not made by this tracker)"),
    ("Production status", "LIVE on pre-longitudinal code. Frontend: Vercel prod dpl_7aCicWiHeqEJuer3grPWqpC3dTow = master 1b39f88 (2026-09-14), function region iad1. Echo is unauthenticated in prod; no scanAnalyses records exist in prod. Gates 01-04 are NOT deployed."),
    ("Current engine version", "Production: scan-0.1.5-provisional. Prepared: scan-0.1.6-provisional (not deployed)."),
    ("Current production engine build", "ECS skin-scan-cluster/skin-scan-service, task definition skin-scan:1 (only revision), image skin-scan:sha-7f3695ef2bfbda37273411766531848774bd098e, digest sha256:668b1f56c440..., 2/2 tasks, /health ok cv2 4.11.0 (verified 2026-10-03)."),
    ("Current deployment state", "Frontend 1b39f88 on Vercel iad1; CV 7f3695e on ECS ap-south-1; Firestore (default) nam5 (US); S3 originals ap-south-1. 0.1.6 release image built locally only (image ID sha256:9a8f96c0a540...), not in ECR."),
    ("Current architecture state", "Longitudinal substrate implemented on branch: server-side raw analysis (gate 01), skinstate-1 computed view (gate 03), trust boundary (gate 02), engine build + native stats inputs (gate 04). No longitudinal comparison, product system, ContextPackage or dashboard exists yet."),
    ("Last tracker update", UPDATED),
    ("Last validated commit", f"{VALIDATED_COMMIT} (typecheck 0 errors, lint clean, vitest 216 passed / 24 skipped, rules emulator 24/24; next build NOT run)"),
]

IMMEDIATE_NEXT = [
    "OWNER: confirm the uncommitted DECISION_LOG.md / DATA_LOCALITY_DECISION.md drafts in the main worktree are superseded by ff08107, then merge infra/india-region (A14).",
    "OWNER: push skin-scan branch engine-build-identity (8dc37dc) to its remote; today it exists on one laptop only (A10).",
    "OWNER: apply dev/prod separation items 1-3 from DEV_PROD_BOUNDARY.md, at minimum unset S3_SKIN_IMAGES_* locally (B01).",
    "OWNER: authorize and execute FIRESTORE_RESET_RUNBOOK.md (destructive; A04). Then verify location asia-south1.",
    "OWNER: authorize 0.1.6 push + task definition skin-scan:2 per ENGINE_0_1_6_DEPLOYMENT_PLAN.md (A12). No Terraform apply.",
    "OWNER: decide D-26 open item (client-asserted captureMethod/persistAttempt in skinstate-1) before comparability work (C02).",
    "OWNER: answer PREFLIGHT G1 (which Google project/tier owns the Gemini key) to unblock H03.",
]

AGENT_NEXT = [
    "H01: read-only check of versioning / Object Lock on projectskin-skin-originals-prod; report erasure IAM needs to owner.",
    "A14 prep: draft a proposed DECISION_LOG reconciliation diff (D-21 duplicate, D-22, D-24) for owner review; do not apply.",
    "C01, E03, F04, H02: write design-gate briefs only; implementation needs owner approval per CLAUDE.md.",
    "I01: draft the production end-to-end verification script for gates 01-04 while A13 is pending (no production execution).",
]

CRITICAL_BLOCKERS = [
    "Production deploy chain is strictly ordered: B01 (items 1-3) -> A04 Firestore reset -> A13 deploy gate 01 together with/before gate 02. Deploying gate 02 without gate 01 makes every Echo request return 409 (GATE_02 section 8).",
    "Developer machine holds production credentials (Firebase Admin key that bypasses rules, S3 PutObject to the production originals bucket, CV secret pointed at production CV, Gemini key); no emulator config. Verified by variable names in .env.local 2026-10-03; DEV_PROD_BOUNDARY.md, PREFLIGHT G3.5.",
    "Production Echo (1b39f88) is unauthenticated and trusts browser-supplied scores; the fix (gate 02) is not deployed.",
    "Gemini data handling unresolved (PREFLIGHT G1): key ownership/tier unknown; prompts carry name and health-related onboarding answers; processing outside India.",
    "Engine 0.1.6 code (8dc37dc) exists only on a local unpushed branch: single-copy loss risk.",
    "No validated measurement noise (SEM/MDC95) exists; every longitudinal change claim is blocked on C04-C07.",
    "DECISION_LOG.md in HEAD has a duplicate D-21 and is missing D-22 and D-24; the approved D-22 lives only on unmerged ff08107.",
]

DISCREPANCIES = [
    ("A04 / D-22 evidence ff08107", "ff08107 is on branch infra/india-region only; it is not an ancestor of longitudinal-platform or master.", "git branch -a --contains ff08107"),
    ("A04 Firestore reset", "Not executed. (default) is still nam5, created 2026-09-08, delete protection and PITR disabled.", "firebase firestore:databases:get (default)"),
    ("A05 Vercel bom1", "vercel.json regions [bom1] exists only in ff08107. Project serverlessFunctionRegion = iad1, functionDefaultRegions = [iad1]; latest prod deployment has regions null.", "vercel api /v9/projects/skin-check"),
    ("A06/A09/B02-B04 'DONE'", "Implemented and tested on branch; none deployed. Production frontend is 1b39f88.", "vercel api; git branch --contains"),
    ("A10 CV commit 8dc37dc", "Exists only on local skin-scan branch engine-build-identity; no remote contains it.", "git -C ../skin-scan branch -a --contains 8dc37dc"),
    ("A12 image digest sha256:9a8f96c0...", "This is the LOCAL image ID (built 2026-10-03, tag sha-8dc37dc...). ECR has only the 7f3695e image. The ECR manifest digest is unknown until push and may differ from the local ID.", "docker images --digests; aws ecr describe-images"),
    ("DECISION_LOG.md numbering", "HEAD: two entries numbered D-21 (311c700 superseded locality draft + gate 01), no D-22, no D-24. Working tree (uncommitted): stale PROPOSED D-22 + D-24. ff08107: APPROVED D-22.", "git show HEAD:docs/architecture/DECISION_LOG.md; git diff"),
    ("D-22 status wording", "Uncommitted working-tree copy says PROPOSED; ff08107 says APPROVED by owner 2026-10-03, reset not executed.", "diff vs git show ff08107:..."),
    ("GATE_03 / GATE_04 docs", "Both state no IMPLEMENTATION_GATE_02.md exists; it now exists (155cb69, 955c1cc). Stale text.", "docs/architecture/IMPLEMENTATION_GATE_0{3,4}.md header"),
    ("Local master", "Local master is 55b6f74; origin/master (deployed) is 1b39f88.", "git rev-parse master origin/master"),
    ("Tracker Markdown tracking", ".gitignore ignores *.md outside whitelisted docs paths; the tracker MD needed an explicit whitelist line (added).", ".gitignore"),
    ("Test coverage note", "npm test skips tests/firestore-rules.test.ts (24); those run only via npm run test:rules (emulator) - run here, 24/24 pass.", "npx vitest run; npm run test:rules"),
    ("Seed statuses", "Seed used non-canonical statuses (DONE LOCALLY, DONE ON BRANCH, NOT STARTED, DESIGNED, PARTIALLY READY, DONE / NEEDS VERIFICATION, LATER / PLANNED). Mapped: Status = implementation state with evidence; production verification recorded in Validation State; NOT STARTED rows derived from dependencies.", "this tracker"),
]

# ---------------------------------------------------------------- master tracker
COLS = ["ID", "Track", "Workstream", "Task", "Description", "Status", "Priority", "Dependency IDs",
        "Gate", "Owner", "Suggested Agent Role", "Repository Area", "Evidence Required",
        "Current Evidence", "Commit", "Validation State", "Blocker", "Next Action",
        "Target Milestone", "Last Updated", "Notes"]

TASKS = []


def t(id, track, ws, task, desc, status, prio, deps="", gate="", owner="Unassigned", role="",
      area="", ev_req="", ev="", commit="", valid="", blocker="", next="", ms="", notes=""):
    TASKS.append([id, track, ws, task, desc, status, prio, deps, gate, owner, role, area, ev_req,
                  ev, commit, valid, blocker, next, ms, UPDATED, notes])


# A - Foundation / Architecture
t("A01", "A", "Foundation / Architecture", "Architecture investigation + repository reconstruction",
  "Reconstruct current frontend/backend/data/deployment architecture; actual vs intended behaviour.",
  "DONE", "P0", role="Systems Architect", area="docs/architecture/",
  ev_req="Committed investigation + baseline", ev="CURRENT_SYSTEM_AUDIT.md, PRODUCTION_BASELINE.md, docs/Longitudinal skin platform architecture.md",
  commit="61c149a, 9078fd6", valid="Doc review. Baseline facts re-verified live 2026-10-03 (Vercel prod 1b39f88/iad1, ECS skin-scan:1, Firestore nam5).",
  next="Update PRODUCTION_BASELINE.md after every production change.", ms="FOUNDATION")
t("A02", "A", "Foundation / Architecture", "Measurement architecture research",
  "Longitudinal measurement architecture, comparability, noise, evidence tiers, validation, digital-shadow framing.",
  "DONE", "P0", role="Measurement Scientist", area="docs/architecture/, docs/research/",
  ev_req="Research documents", ev="LONGITUDINAL_METHOD.md, SYSTEM_MODEL.md, EVALUATION_PLAN.md, MEASUREMENT_VALIDATION_STATUS.md",
  commit="61c149a, 59d9923", valid="Research synthesis only; not empirical validation.",
  next="None; superseded in part by C03 protocol (D-24).", ms="FOUNDATION")
t("A03", "A", "Foundation / Architecture", "Production preflight G1-G4",
  "Security, data locality, privacy, persistence, recoverability, production risks.",
  "DONE", "P0", role="Security/Privacy Reviewer", area="docs/architecture/PREFLIGHT_AUDIT.md",
  ev_req="Audit report", ev="PREFLIGHT_AUDIT.md", commit="9078fd6",
  valid="Read-only audit. Open findings tracked as B01 (G3.5), A04 (G4.2), H03 (G1).",
  next="None; follow-ups live in B01, A04, H03.", ms="FOUNDATION")
t("A04", "A", "Foundation / Architecture", "Firestore India-region reset (D-22)",
  "Recreate disposable test (default) Firestore in asia-south1, no data migration; Firebase Auth untouched.",
  "OWNER ACTION", "P0", deps="A14, B01", gate="D-22", owner="Owner", role="Backend/Data Engineer",
  area="Firebase project projectskin-a202c; docs/architecture/FIRESTORE_RESET_RUNBOOK.md (on ff08107)",
  ev_req="Owner authorization; post-reset `firebase firestore:databases:get` showing asia-south1, delete protection, PITR, backups",
  ev="Runbook + decision on infra/india-region (ff08107)", commit="ff08107 (unmerged)",
  valid="Live 2026-10-03: (default) location nam5, delete protection disabled, PITR disabled. Reset NOT executed.",
  blocker="Destructive production action: requires explicit owner authorization and execution. Recommended after B01 items 1-3 so laptops do not write into the new database.",
  next="Owner reviews FIRESTORE_RESET_RUNBOOK.md, authorizes, executes; record new location/settings here.", ms="FOUNDATION")
t("A05", "A", "Foundation / Architecture", "Vercel Mumbai region (bom1)",
  "Run server functions in bom1; verify project/dashboard config too.",
  "PREPARED", "P1", deps="A04, A14", gate="D-22", owner="Owner", role="Backend/Data Engineer",
  area="vercel.json; Vercel project skin-check",
  ev_req="Prod deployment inspect showing bom1 function region", ev="vercel.json regions [bom1] on ff08107",
  commit="ff08107 (unmerged)",
  valid="Live 2026-10-03: project serverlessFunctionRegion iad1, functionDefaultRegions [iad1]; prod deployment regions null. Not on bom1.",
  next="Merge ff08107 into the release path; ship in the same window as A04; verify with `vercel inspect`; consider setting project default region too.", ms="FOUNDATION")
t("A06", "A", "Foundation / Architecture", "Gate 01 durable persistence/recovery",
  "Preserve complete raw analysis server-side (scanAnalyses); durable, recoverable save.",
  "DONE", "P0", gate="Gate 01 / D-21", role="Backend/Data Engineer",
  area="app/api/skin-scan, lib/scan-analysis-record.ts", ev_req="Commit + gate report + tests",
  ev="IMPLEMENTATION_GATE_01.md; tests/skin-scan-route-persist, scan-recovery, scan-analysis-record pass", commit="6e9081d, 428c63d, 311c700",
  valid=NOT_DEPLOYED, next="Deploy via A13 after A04.", ms="TRUSTED MEASUREMENT")
t("A07", "A", "Foundation / Architecture", "SkinState-1 contract",
  "Freeze versioned skinstate-1 contract.", "DONE", "P0", gate="D-23", role="Systems Architect",
  area="docs/architecture/SKIN_STATE_CONTRACT_FINAL.md", ev_req="Committed frozen contract",
  ev="SKIN_STATE_CONTRACT_FINAL.md + Amendment A1", commit="b049685",
  valid="Contract frozen. Open owner decision in D-26 on client-asserted captureMethod/persistAttempt; contract text 'no client sends captureMethod' is stale since gate 04.",
  next="Owner decides D-26 open item (may require skinstate-2).", ms="TRUSTED MEASUREMENT")
t("A08", "A", "Foundation / Architecture", "SkinState-1 computed view",
  "Compute SkinState from trusted persisted scanAnalyses; not stored.", "DONE", "P0", deps="A06, A07",
  gate="Gate 03 / D-23", role="Backend/Data Engineer", area="lib/skin-state*", ev_req="Commit + tests",
  ev="IMPLEMENTATION_GATE_03.md; tests/skin-state-loader pass", commit="dd9a3fe",
  valid="Pure library; loadTrustedSkinState has no production caller (GATE_02 section 7.8). " + NOT_DEPLOYED,
  next="First consumer arrives with F01/G01.", ms="TRUSTED MEASUREMENT")
t("A09", "A", "Foundation / Architecture", "Gate 02 trust / integrity",
  "Trusted analysis context, ownership, server-authoritative state, Echo isolation.", "DONE", "P0",
  deps="A06", gate="Gate 02 / D-26", role="Security/Privacy Reviewer",
  area="app/api/skin-check/chat, lib/echo-trusted-context.ts", ev_req="Commit + tests + live E2E in production",
  ev="IMPLEMENTATION_GATE_02.md; chat-auth tests; mutation testing; local probe reached Gemini call (dummy key)",
  commit="155cb69, 955c1cc",
  valid="Closed on branch. " + NOT_DEPLOYED + " Production trust NOT established until I01 end-to-end check.",
  next="Deploy only with/after gate 01 (A13); then verify live (I01).", ms="TRUSTED MEASUREMENT",
  notes="15 residual risks listed in GATE_02 section 7.")
t("A10", "A", "Foundation / Architecture", "Engine 0.1.6 metadata/input support",
  "engine_build, image tag, engine version, stats=both, capture method persistence.", "DONE", "P1",
  gate="Gate 04 / D-07", role="CV/ML Engineer", area="skin-scan src/app, version.py; skin-check app/api/skin-scan",
  ev_req="Commits in both repos + tests", ev="IMPLEMENTATION_GATE_04.md", commit="skin-scan 8dc37dc; skin-check 1e4919d",
  valid="Local only. 8dc37dc is on local skin-scan branch engine-build-identity, not pushed to any remote. 1e4919d not deployed.",
  next="Owner pushes engine-build-identity to the skin-scan remote.", ms="TRUSTED MEASUREMENT")
t("A11", "A", "Foundation / Architecture", "Engine 0.1.6 Phase 0 replay",
  "30 deterministic replays vs 0.1.5 on the release-candidate image.", "DONE", "P1", deps="A10",
  role="Measurement Scientist", area="docs/research/measurement-validation/PHASE_0_0_1_6_RESULTS.md",
  ev_req="Pre-registered replay report", ev="PHASE_0_0_1_6_RESULTS.md; addendum ebbe755", commit="91d20bc, ebbe755, 631a9fa",
  valid="Screening, not validation. 30/30 bit-identical to 0.1.5 apart from version/engine_build; 500/500 deterministic on image sha256:9a8f96c0.... One face, no device encoders, no ColorChecker, no native amd64.",
  next="None; no architecture decision changed.", ms="CALIBRATION")
t("A12", "A", "Foundation / Architecture", "Engine 0.1.6 production deployment",
  "Exact-image deploy: push, verify ECR digest, task def pinned by digest, update service; rollback to skin-scan:1. No Terraform apply.",
  "PREPARED", "P1", deps="A10, A11", gate="ENGINE_0_1_6_DEPLOYMENT_PLAN", owner="Owner", role="CV/ML Engineer",
  area="ECR skin-scan; ECS skin-scan-service", ev_req="ECR digest, task def skin-scan:2 pinned by digest, /health/ready engine_build check, rollback tested or documented",
  ev="ENGINE_0_1_6_DEPLOYMENT_PLAN.md; local release image built and replayed (91d20bc)", commit="9808d91",
  valid="Live 2026-10-03: service on skin-scan:1, image sha-7f3695e (digest 668b1f56...); ECR has no sha-8dc37dc tag.",
  next="Owner authorizes; execute plan from 'push' onward; record ECR manifest digest; register skin-scan:2 by digest; update service.",
  ms="TRUSTED MEASUREMENT", notes="Activates D-26 residual (client-asserted captureMethod in comparability key) once frontend gate 04 is also deployed.")
t("A13", "A", "Foundation / Architecture", "Frontend gates 01-04 production deployment",
  "Merge longitudinal-platform to master and deploy gates 01, 02, 03, 04 in the required order.",
  "BLOCKED", "P0", deps="A04, A05, B01", gate="GATE_02 section 8", owner="Owner", role="Backend/Data Engineer",
  area="branch longitudinal-platform -> master; Vercel", ev_req="Merged PR, prod deployment ID + commit, post-deploy smoke (scan saves scanAnalyses; Echo authenticated)",
  ev="", valid="Not started. Added by tracker 2026-10-03 (missing from seed inventory).",
  blocker="A04 Firestore reset not executed; gate 02 cannot deploy without gate 01 (Echo 409s).",
  next="After A04: open PR longitudinal-platform -> master, preview deploy, smoke test, promote.", ms="TRUSTED MEASUREMENT")
t("A14", "A", "Foundation / Architecture", "Reconcile DECISION_LOG and merge infra/india-region",
  "Fix duplicate D-21, missing D-22/D-24 in HEAD; bring approved D-22 (ff08107) into the main line.",
  "OWNER ACTION", "P0", owner="Owner", role="Program Lead", area="docs/architecture/DECISION_LOG.md",
  ev_req="Commit with unique, complete D-numbering", ev="Discrepancy recorded in this tracker",
  valid="Working tree holds uncommitted stale drafts not made by this tracker.",
  blocker="Uncommitted working-tree edits belong to another session; owner must confirm they are superseded before they are discarded or merged.",
  next="Owner confirms; merge infra/india-region; commit D-24; renumber/annotate the superseded D-21 draft.", ms="FOUNDATION",
  notes="Added by tracker 2026-10-03.")

# B - Trust / Security
t("B01", "B", "Trust / Security", "Dev/prod separation",
  "Stop local dev writing production resources; emulators / isolated dev resources.", "OWNER ACTION", "P0",
  owner="Owner", role="Security/Privacy Reviewer", area=".env.local, lib/firebase/client.ts, lib/firebase/admin.ts",
  ev_req="Admin key removed from laptops + rotated; S3 vars unset or dev bucket; emulator switch committed",
  ev="DEV_PROD_BOUNDARY.md (on ff08107)", commit="ff08107 (unmerged)",
  valid="Not started. Verified 2026-10-03 (variable names only, no values read): .env.local defines FIREBASE_PRIVATE_KEY/FIREBASE_CLIENT_EMAIL, S3_SKIN_IMAGES_ACCESS_KEY_ID/SECRET, SCAN_API_SECRET with SKIN_SCAN_API_URL=https://api.theprojectskin.com, GEMINI_API_KEY; no emulator variables.",
  blocker="Credential removal/rotation and dev resources need the owner; emulator code switch needs an approved gate.",
  next="Owner: unset S3_SKIN_IMAGES_* locally today; approve dev-only emulator gate; rotate Admin key after removal.", ms="PRIVACY/SAFETY")
t("B02", "B", "Trust / Security", "Server-authoritative analysis writes",
  "Browser is not authoritative for analysis-derived scores.", "DONE", "P0", deps="A06", gate="D-02 / D-26",
  role="Backend/Data Engineer", area="app/api/skin-scan, lib/skin-state*", ev_req="Commit + tests",
  ev="scanAnalyses written server-side; SkinState trust is path-structural", commit="6e9081d, 155cb69",
  valid=NOT_DEPLOYED + " Residual: dashboard still renders browser-held result (GATE_02 7.3); skinChecks.result remains owner-writable but no consumer reads it.",
  next="Deploy via A13; dashboard read path moves to server state in G01.", ms="TRUSTED MEASUREMENT")
t("B03", "B", "Trust / Security", "Echo authentication and ownership",
  "Echo derives identity/context server-side and verifies ownership.", "DONE", "P0", deps="A06", gate="Gate 02 / D-26",
  role="Security/Privacy Reviewer", area="app/api/skin-check/chat", ev_req="Commit + tests + live E2E",
  ev="skin-check-chat-auth tests (13); cross-user 409", commit="155cb69, 955c1cc",
  valid=NOT_DEPLOYED + " Production Echo is unauthenticated today.", next="Deploy via A13; verify live in I01.", ms="TRUSTED MEASUREMENT")
t("B04", "B", "Trust / Security", "Gemini browser-state isolation",
  "Gemini never trusts browser-supplied scores/context.", "DONE", "P0", deps="A06", gate="Gate 02 / D-26",
  role="Gemini/Context Engineer", area="lib/echo-trusted-context.ts, lib/gemini.ts", ev_req="Commit + tests + live E2E",
  ev="Strict body schema rejects scores/uid; HMAC-signed history", commit="155cb69, 955c1cc",
  valid=NOT_DEPLOYED + " Residual: user's own name still in system instruction (GATE_02 7.12).",
  next="Deploy via A13; verify live in I01.", ms="GEMINI CONTEXT LAYER")

# C - Measurement quality / validation
t("C01", "C", "Measurement Quality", "Guided capture quality gate",
  "On-device checks: pose, expression, exposure, distance, blur, framing; retake loop.", "UNBLOCKED", "P1",
  role="Frontend/Dashboard Engineer", area="components/skin-check capture step",
  ev_req="Design doc + approved gate + tests", valid="Not started.",
  next="Write design gate: checks, thresholds, retake UX; exposure first (dominant Phase 0 risk).", ms="CALIBRATION")
t("C02", "C", "Measurement Quality", "Per-region / per-metric comparability",
  "Decide whether two scans are comparable before computing change.", "BLOCKED", "P0", deps="A07, A12, A13, C07",
  gate="D-13", role="Measurement Scientist", area="lib/skin-state*", ev_req="Comparability function + tests + noise inputs",
  valid="Not started. engine.build is null in all records until 0.1.6 + gate 04 are deployed.",
  blocker="engine.build not populated (A12, A13); D-26 owner decision on client-asserted captureMethod; per-metric thresholds need C07.",
  next="Owner decides D-26; implement key-based comparability after A12/A13; noise thresholds after C07.", ms="LONGITUDINAL INTELLIGENCE")
t("C03", "C", "Measurement Quality", "Human measurement validation protocol",
  "Protocol for human validation and repeatability.", "DONE", "P0", gate="D-24 (uncommitted)",
  role="Measurement Scientist", area="docs/research/measurement-validation/", ev_req="Committed protocol",
  ev="PROTOCOL.md, ANALYSIS_PLAN.md, ACCEPTANCE_CRITERIA.md, DATA_SPEC.md", commit="59d9923",
  valid="Design only; not executed. D-24 adopting it exists only as uncommitted working-tree text.",
  next="Commit D-24 (A14); owner may revise MD-1 endpoint cut before Phase 1 pre-registration (D-25).", ms="CALIBRATION")
t("C04", "C", "Measurement Quality", "Human validation pilot",
  "Run pilot of C03 protocol.", "BLOCKED", "P0", deps="C03, A12, A13, H01", role="Measurement Scientist",
  area="docs/research/measurement-validation/", ev_req="Pre-registration + pilot dataset + report",
  valid="Not started.",
  blocker="Protocol scores via the production HTTP path (D-24), which needs 0.1.6 + gate 04 deployed (A12, A13); consent (H01); owner approval of recruitment, devices, budget, ColorChecker/colorimeter.",
  next="Owner approves pilot logistics; pre-register Phase 1.", ms="CALIBRATION")
t("C05", "C", "Measurement Quality", "Full test-retest / calibration",
  "D0/D7/D28 test-retest per field x region x tone band.", "RESEARCH", "P0", deps="C04",
  role="Measurement Scientist", ev_req="Pre-registered study report", valid="Not started.",
  blocker="Waits on C04 pilot.", next="Size from pilot variance.", ms="CALIBRATION")
t("C06", "C", "Measurement Quality", "Tone-stratified validation / model card",
  "Performance/noise across MST bands 1-4/5-6/7-10; documented limitations.", "RESEARCH", "P0", deps="C05",
  role="Measurement Scientist", ev_req="Per-band results + model card", valid="Not started.",
  blocker="Waits on C05 data.", next="Define model card template alongside C05.", ms="CALIBRATION")
t("C07", "C", "Measurement Quality", "Validated measurement/noise tables",
  "Empirical reliability, SEM/MDC95, metric/region conditions.", "RESEARCH", "P0", deps="C05, C06",
  role="Measurement Scientist", ev_req="Versioned noise table tied to engine build", valid="Not started.",
  blocker="Waits on C05/C06.", next="Fix table schema (field x region x band x engine build).", ms="CALIBRATION",
  notes="Provisional candidates only: native L* (cheeks, forehead), native a* (cheeks, nose). Not validated.")

# D - Longitudinal intelligence
t("D01", "D", "Longitudinal Intelligence", "Baseline window definition",
  "How a user's baseline is formed.", "RESEARCH", "P1", deps="C07", role="Measurement Scientist",
  ev_req="Method doc + decision", valid="Not started.", blocker="Needs between-day noise (C07).",
  next="Draft candidate baseline rules against LONGITUDINAL_METHOD.md.", ms="LONGITUDINAL INTELLIGENCE")
t("D02", "D", "Longitudinal Intelligence", "Change-vs-noise engine",
  "Report change only beyond comparability + MDC rules. Never naive % change.", "BLOCKED", "P0", deps="C02, C07, D01",
  gate="D-13 / D-24", role="Backend/Data Engineer", ev_req="Engine + tests + noise table version",
  valid="Not started.", blocker="Needs comparability (C02) and validated noise (C07).",
  next="Wait; design pairwise interface first (D-13).", ms="LONGITUDINAL INTELLIGENCE")
t("D03", "D", "Longitudinal Intelligence", "Persistence/evidence tiers",
  "Separate transient observations from persistent patterns.", "BLOCKED", "P1", deps="D02",
  role="Measurement Scientist", ev_req="Tier definitions + tests", valid="Not started.",
  blocker="Needs D02.", next="Wait on D02.", ms="LONGITUDINAL INTELLIGENCE",
  notes="D-12/D-13 reject research-doc tiers 3-4 and the 4/10-point ladder as primary output.")
t("D04", "D", "Longitudinal Intelligence", "Scheduled anchor scans / irregular sampling",
  "Sampling strategy.", "LATER", "P3", deps="D02", next="Revisit after D02.", ms="LONGITUDINAL INTELLIGENCE")
t("D05", "D", "Longitudinal Intelligence", "Engine-version re-baselining",
  "Handle engine build changes across history.", "LATER", "P3", deps="D02",
  next="Revisit at first engine change after D02.", ms="LONGITUDINAL INTELLIGENCE")

# E - Products / routine / adherence
t("E01", "E", "Products / Routine", "Product catalog normalization",
  "Versioned own-brand catalog: ingredients, classes, claims, constraints, cost, availability.", "OWNER ACTION", "P1",
  gate="D-08 / D-09", owner="Owner", role="Product/Routine Rules Engineer", ev_req="Owner-supplied master data + versioned schema",
  valid="Not started.", blocker="Product master data and claims must come from the owner; D-09 OPEN-R1 (do product cards return).",
  next="Owner supplies catalog data and decides OPEN-R1.", ms="ROUTINE/PRODUCT SYSTEM")
t("E02", "E", "Products / Routine", "Product eligibility / safety rules",
  "Deterministic, versioned eligibility + safety constraints.", "BLOCKED", "P1", deps="E01, E05", gate="D-08",
  role="Product/Routine Rules Engineer", ev_req="Pure function + version constant + tests", valid="Not started.",
  blocker="Needs catalog (E01) and red-flag criteria (E05).", next="Wait on E01.", ms="ROUTINE/PRODUCT SYSTEM")
t("E03", "E", "Products / Routine", "Routine model + event history",
  "Products, routines, starts, stops, substitutions, schedule history (append-only).", "UNBLOCKED", "P2", deps="A04",
  role="Backend/Data Engineer", ev_req="Data model doc + approved gate + tests", valid="Not started.",
  next="Design data model on branch; no production writes before A04.", ms="ROUTINE/PRODUCT SYSTEM",
  notes="A04 dependency applies to deployment only.")
t("E04", "E", "Products / Routine", "Adherence summary",
  "used / skipped / unknown; missing log is never 'skipped'.", "BLOCKED", "P2", deps="E03",
  role="Product/Routine Rules Engineer", ev_req="Function + tests", valid="Not started.",
  blocker="Needs E03.", next="Wait on E03.", ms="ROUTINE/PRODUCT SYSTEM")
t("E05", "E", "Products / Routine", "Red-flag triage interlock",
  "Safety layer overrides cosmetic recommendations.", "OWNER ACTION", "P0", gate="D-14", owner="Owner",
  role="Product/Routine Rules Engineer", ev_req="Approved red-flag criteria + tests (EN + Hinglish)", valid="Not started.",
  blocker="Red-flag criteria need qualified clinical sign-off sourced by the owner.",
  next="Owner nominates clinical reviewer; engineer drafts criteria list for review.", ms="PRIVACY/SAFETY")

# F - Context / Gemini
t("F01", "F", "Context / Gemini", "Deterministic ContextPackage",
  "Trusted context: SkinState, history, routines, adherence, uncertainty, evidence tier, safety.", "BLOCKED", "P1",
  deps="A08, D02, E03, E04, E05", gate="D-14", role="Gemini/Context Engineer", ev_req="Pure builder + tests",
  ev="Precursor: lib/echo-trusted-context.ts (155cb69)", valid="Not started.",
  blocker="Needs longitudinal (D02), routine (E03/E04) and safety (E05) inputs.",
  next="Wait; may define schema with null sections earlier if approved.", ms="GEMINI CONTEXT LAYER")
t("F02", "F", "Context / Gemini", "Gemini schema renderer",
  "Gemini as explanation layer over a JSON schema with fact-id enum.", "BLOCKED", "P1", deps="F01", gate="D-14",
  role="Gemini/Context Engineer", area="lib/gemini.ts", ev_req="Schema + renderer + tests",
  ev="Precursor: GEMINI_BOUNDARY.md; server-grounded Echo prompt (155cb69)", valid="Schema renderer not built.",
  blocker="Needs F01.", next="Wait on F01.", ms="GEMINI CONTEXT LAYER", notes="Seed status PARTIALLY READY.")
t("F03", "F", "Context / Gemini", "Grounding / number / reference validator",
  "Post-generation validator (digits, comparative/efficacy lexicon, fallback).", "BLOCKED", "P1", deps="F02",
  gate="D-14 / D-12", role="Gemini/Context Engineer", ev_req="Validator + adversarial tests", valid="Not started.",
  blocker="Needs F02 output schema.", next="Wait on F02.", ms="GEMINI CONTEXT LAYER")
t("F04", "F", "Context / Gemini", "UserAssertion quarantine",
  "User claims never silently become measured state.", "UNBLOCKED", "P2", gate="D-14 (store rejected for now)",
  role="Gemini/Context Engineer", ev_req="Design + tests", ev="Answers already labelled SELF_REPORTED (GATE_02 7.1)",
  valid="Not started.", next="Design quarantine labelling; D-14 lists the UserAssertion store as LATER, so scope to labelling only.",
  ms="GEMINI CONTEXT LAYER")

# G - Dashboard
t("G01", "G", "Dashboard", "Longitudinal dashboard",
  "Latest state, history, comparable change only, confidence, routine, adherence, evidence, privacy controls.",
  "BLOCKED", "P1", deps="A13, C02, D02", role="Frontend/Dashboard Engineer", ev_req="Approved design + tests + E2E",
  valid="Not started.", blocker="Comparable-change views need C02/D02; server-read history needs gate 01 in prod (A13).",
  next="Latest-state/history view can be designed after A13.", ms="DASHBOARD")
t("G02", "G", "Dashboard", "Product-response timeline UX",
  "Routine events next to observations, no causal claims.", "BLOCKED", "P2", deps="E03, G01", gate="D-12 (LOCKED)",
  role="Frontend/Dashboard Engineer", ev_req="Design reviewed against D-12", valid="Not started.",
  blocker="Needs E03 and G01.", next="Wait.", ms="DASHBOARD")

# H - Privacy / safety
t("H01", "H", "Privacy / Safety", "Consent / retention / deletion fan-out",
  "Server-written consent ledger; erasure S3 -> Firestore -> Auth.", "BLOCKED", "P0", deps="A04", gate="D-16 / D-17",
  role="Security/Privacy Reviewer", ev_req="Ledger + erasure route + tests + live verification", valid="Not started.",
  blocker="D-17 blocked on IAM (S3 PutObject-only) and unknown bucket Object Lock/versioning.",
  next="Inspect bucket versioning/Object Lock (read-only); propose erasure IAM to owner.", ms="PRIVACY/SAFETY")
t("H02", "H", "Privacy / Safety", "18+ gate / safety UX",
  "Server-side 18+ attestation; re-confirm sensitive answers per scan.", "UNBLOCKED", "P1", gate="D-19",
  role="Frontend/Dashboard Engineer", ev_req="Gate + tests", valid="Not started; 18+ check is client-only today.",
  next="Write gate brief for D-19.", ms="PRIVACY/SAFETY")
t("H03", "H", "Privacy / Safety", "In-region prompt/data audit",
  "What reaches Gemini, where processed/logged; consent and minimization.", "OWNER ACTION", "P0", gate="PREFLIGHT G1 / D-14",
  owner="Owner", role="Security/Privacy Reviewer", ev_req="Field inventory + key ownership/tier + consent copy",
  valid="Not started.", blocker="Owner must state which Google project and tier own the production Gemini key (paid tier is a D-14 release gate).",
  next="Owner answers G1; reviewer inventories prompt fields (name, onboarding answers).", ms="PRIVACY/SAFETY")

# I - Release
t("I01", "I", "Release Validation", "End-to-end production evaluation",
  "Verify trusted flow, persistence, Echo isolation in production.", "BLOCKED", "P0",
  deps="A12, A13, B01, H01, H03", role="QA/Test Engineer", ev_req="Recorded E2E run against prod", valid="Not started.",
  blocker="Nothing longitudinal is deployed yet.", next="Write E2E script while A13 is pending.", ms="RELEASE")
t("I02", "I", "Release Validation", "Full red-team sweep",
  "Adversarial review across trust, Gemini, privacy, safety.", "BLOCKED", "P0", deps="I01, F03, E05",
  role="Red-Team Reviewer", ev_req="Red-team report", valid="Not started.", blocker="Needs I01 and the Gemini/safety layers.",
  next="Wait.", ms="RED TEAM")

# J - Later
for jid, task in [("J01", "Population priors / mixed effects"), ("J02", "ML recommendation ranking"),
                  ("J03", "Indian cross-user benchmark"), ("J04", "Forecasting"), ("J05", "True digital twin")]:
    t(jid, "J", "Later / Research-dependent", task, "Not in core release.", "LATER", "P4",
      next="Do not start without evidence and explicit approval.", ms="RELEASE",
      notes="Forecasting needs explicit approval (CLAUDE.md, D-12)." if jid == "J04" else "")

# ---------------------------------------------------------------- research track
RESEARCH_COLS = ["Research ID", "Research question", "Why it matters", "Evidence currently available",
                 "Evidence missing", "Candidate metrics/signals", "Protocol required", "Decision enabled",
                 "Current decision", "Confidence", "Source/document", "Owner", "Status", "Next experiment"]
R = [
    ["R01", "Test-retest reliability of native colour fields", "No change claim without it", "Phase 0 synthetic screening, deterministic engine", "Human between-day repeats", "native L* (cheeks, forehead), native a* (cheeks, nose)", "C03 D0/D7/D28", "C05, D02", "MD-1 / D-25: L*, a* primary endpoints", "Low (screening only)", "PROTOCOL.md, PHASE_0_RESULTS.md", "Measurement Scientist", "RESEARCH", "C04 pilot"],
    ["R02", "SEM / MDC95 per field x region x band", "Sets change thresholds", "None empirical", "Human data", "as R01", "C03", "D02 thresholds", "Uncertainty = unknown (gate 03)", "None", "ANALYSIS_PLAN.md, ACCEPTANCE_CRITERIA.md", "Measurement Scientist", "RESEARCH", "From C05"],
    ["R03", "Inter-device comparability", "Users change phones", "None", "Real device encoders", "as R01", "Own-phone arm of C03", "C02 key (captureMethod)", "captureMethod in comparability key (D-13)", "None", "PROTOCOL.md", "Measurement Scientist", "RESEARCH", "Device arm in pilot"],
    ["R04", "Tone-stratified performance", "Fails in any band = fails overall", "None", "MST-quota sample", "as R01", "Bands 1-4/5-6/7-10, ~160 enrolled", "C06 model card", "D-24 (uncommitted)", "None", "D-24, PROTOCOL.md", "Measurement Scientist", "RESEARCH", "Quota recruitment plan"],
    ["R05", "Capture quality effect", "Exposure is dominant Phase 0 risk", "Phase 0 exposure sensitivity", "Real-world capture variance", "exposure, blur, pose", "C01 design + pilot", "C01 thresholds", "None", "Low", "PHASE_0_RESULTS.md", "CV/ML Engineer", "RESEARCH", "Exposure sweep with real devices"],
    ["R06", "Longitudinal baseline strategy", "Defines 'change from what'", "LONGITUDINAL_METHOD.md", "Noise data, user scan cadence", "-", "-", "D01", "Pairwise first (D-13)", "Low", "LONGITUDINAL_METHOD.md", "Measurement Scientist", "RESEARCH", "After C05"],
    ["R07", "Evidence tiers", "Prevents over-claiming", "Research doc tiers", "Noise + persistence data", "-", "-", "D03", "Tiers 3-4 not adopted (D-12)", "Medium (policy)", "DECISION_LOG D-12, D-13", "Measurement Scientist", "RESEARCH", "After D02"],
    ["R08", "Product-response attribution", "Users expect it", "None", "Everything", "-", "-", "None in core", "LOCKED: no attribution/efficacy (D-12)", "High (policy)", "CLAUDE.md, D-12", "Owner", "LATER", "None without explicit approval"],
    ["R09", "Gemini data minimization", "Health data to a US processor", "GATE_02 field restrictions; name still sent", "Key ownership/tier, logging, region", "-", "Field inventory", "H03, D-14 release gate", "Paid tier is a release gate (D-14)", "Low", "PREFLIGHT_AUDIT.md G1, GEMINI_BOUNDARY.md", "Owner", "OWNER ACTION", "Owner answers G1"],
    ["R10", "Privacy / data locality", "India-first", "Firestore nam5 verified; ECS/S3 ap-south-1; Vercel iad1", "Post-reset verification", "-", "Runbook", "A04, A05", "D-22 approved (ff08107)", "High", "DATA_LOCALITY_DECISION.md", "Owner", "OWNER ACTION", "Execute A04"],
    ["R11", "Population priors", "Small per-user data", "None", "Validated per-user noise", "-", "-", "J01", "LATER", "None", "-", "Measurement Scientist", "LATER", "None"],
    ["R12", "Forecasting feasibility", "Product ask", "None", "Everything", "-", "-", "J04", "Not approved (CLAUDE.md)", "None", "CLAUDE.md, D-12", "Owner", "LATER", "None without explicit approval"],
]

# ---------------------------------------------------------------- Claude operations
OPS_COLS = ["Agent/session name", "Role", "Current task", "Parent task", "Worktree", "Branch", "Status",
            "Dependencies", "Files touched", "Tests run", "Findings", "Messages sent", "Decision impact",
            "Commit", "Review status", "Shutdown state"]
OPS = [
    ["tracker-session-2026-10-03", "Program Lead", "Create execution tracker", "-", "main (skin-check)", "longitudinal-platform", "DONE", "-",
     "scripts/build_execution_tracker.py, docs/PROJECT_SKIN_EXECUTION_TRACKER.{md,xlsx}, .gitignore", "typecheck, lint, vitest, test:rules",
     "See discrepancies", "-", "None (no decision changed)", "uncommitted", "Awaiting owner review", "Active"],
    ["india-region session", "Backend/Data Engineer", "Prepare India-region Firestore + Vercel", "A04, A05", ".claude/worktrees/india-region", "infra/india-region", "DONE", "-",
     "vercel.json, DATA_LOCALITY_DECISION.md, DECISION_LOG.md, DEV_PROD_BOUNDARY.md, FIRESTORE_RESET_RUNBOOK.md", "unknown", "D-22", "unknown", "D-22 approved", "ff08107", "Not merged", "Worktree still exists (clean)"],
    ["earlier sessions (gates 01-04, Phase 0)", "unknown", "-", "A01-A12, B02-B04", "unknown", "longitudinal-platform", "unknown", "-", "see commits", "see gate reports", "see gate reports", "unknown", "see DECISION_LOG", "61c149a..91d20bc", "unknown", "unknown - not reconstructable from repo"],
]
ROLES = [
    ("Program Lead", "Owns this tracker, sequencing, evidence checks"),
    ("Systems Architect", "Contracts, data model, DECISION_LOG"),
    ("Measurement Scientist", "Protocols, noise, comparability, evidence tiers"),
    ("CV/ML Engineer", "skin-scan engine, images, deployments"),
    ("Backend/Data Engineer", "API routes, Firestore, persistence, infra"),
    ("Frontend/Dashboard Engineer", "Capture UX, dashboard"),
    ("Security/Privacy Reviewer", "Trust boundary, credentials, consent, erasure"),
    ("Research Reviewer", "Independent review of research claims"),
    ("QA/Test Engineer", "Tests, E2E, regression"),
    ("Product/Routine Rules Engineer", "Catalog, eligibility, adherence, red flags"),
    ("Gemini/Context Engineer", "ContextPackage, renderer, validator"),
    ("Red-Team Reviewer", "Adversarial sweeps"),
]

# ---------------------------------------------------------------- decisions
DEC_COLS = ["Decision ID", "Date", "Decision", "Context", "Alternatives considered", "Chosen direction",
            "Why", "Evidence", "Affected tasks", "Reopening conditions", "Owner", "Status"]
DECISIONS = [
    ["D-22", "2026-10-03", "Firestore region", "PREFLIGHT G4.2; prod data is test data", "Named DB; new project; stay nam5", "Recreate disposable (default) in asia-south1, no migration; Auth untouched; Vercel bom1", "India-first, cost, location immutable", "ff08107 (unmerged); absent from DECISION_LOG in HEAD", "A04, A05, A13", "asia-south1 refused for (default) -> empty named DB", "Owner", "PREPARED / OWNER ACTION PENDING"],
    ["D-23", "2026-10-03", "SkinState contract", "Gate 01 record exists", "Stored skinStates collection", "Frozen skinstate-1, computed view over scanAnalyses", "One source of truth", "b049685, SKIN_STATE_CONTRACT_FINAL.md", "A07, A08", "Any mapping change -> new version", "Systems Architect", "DONE"],
    ["D-02", "2026-10-03", "Server-authoritative analysis", "Client-written result was forgeable", "Rules validation; S3 sidecar", "Server writes the measurement", "Rules cannot prove CV origin", "6e9081d, D-26", "A06, B02", "-", "Systems Architect", "DONE (branch; not deployed)"],
    ["D-23 (clause: legacy)", "2026-10-03", "Legacy scans remain legacy", "Legacy skinChecks.result", "Migrate/upgrade", "legacy_display only; never upgraded", "No raw data or provenance", "SKIN_STATE_CONTRACT_FINAL.md; GATE_03", "A08, G01", "-", "Systems Architect", "DONE"],
    ["TD-01", "-", "Malformed raw analysis is rejected", "Raw records may be incomplete", "Best-effort mapping", "toSkinState returns invalid (11 malformed cases tested)", "No fabricated values", "IMPLEMENTATION_GATE_03.md (dd9a3fe)", "A08", "-", "Systems Architect", "DONE (not in DECISION_LOG)"],
    ["TD-02", "-", "Uncertainty unknown until calibrated", "No calibration exists", "Heuristic intervals", "uncertainty.policy = unknown always", "No unvalidated precision", "GATE_03 (dd9a3fe)", "A08, C07, D02", "C07 delivers calibration", "Measurement Scientist", "DONE (not in DECISION_LOG)"],
    ["TD-03", "-", "Empty regions = region absent", "Engine writes 0.0 for empty masks", "Expose 0.0", "region_px 0 -> absent_region, nulls", "0.0 indistinguishable from real zero", "SKIN_STATE_CONTRACT_FINAL.md, GATE_03", "A08", "-", "Systems Architect", "DONE (not in DECISION_LOG)"],
    ["D-25", "2026-10-03", "Native L*/a* are provisional research candidates only", "Phase 0 screening", "Keep b*, ITA, lbpEntropy", "Primary endpoints for the human study only; not validated measurements", "Only survivors of screening", "MEASUREMENT_METRIC_DECISION.md, 59d9923, 91d20bc", "C03-C07", "Owner may reverse before Phase 1 pre-registration", "Owner", "ACCEPTED for study design"],
    ["D-13 / D-24", "-", "Longitudinal comparisons require comparability + noise/evidence", "Naive % change misleads", "Research-doc point ladder", "Pairwise, keyed comparability; change claims gated on validation protocol", "Measurement noise unknown", "DECISION_LOG D-13; D-24 uncommitted", "C02, D02, G01", "-", "Measurement Scientist", "PROPOSED (D-24 uncommitted)"],
    ["D-08", "-", "Recommender deterministic and versioned", "4 SKUs, ~6 rules", "DMN/rules engine; ML ranking", "Plain TS pure function + version constants", "Auditable", "DECISION_LOG D-08", "E02, J02", ">20 SKUs or non-engineer rule authors", "Systems Architect", "PROPOSED"],
    ["D-14", "-", "Gemini is explanation layer, not authority", "Gemini hallucination/claims risk", "Gemini owns context", "Server-built ContextPackage, pre-filters, schema, validator", "CLAUDE.md", "DECISION_LOG D-14; GATE_02", "B04, F01-F04", "-", "Systems Architect", "PROPOSED (paid tier is release gate)"],
    ["D-12", "-", "No attribution/efficacy/forecasting", "CLAUDE.md", "Evidence tiers 3-4", "Prohibited in UI, Gemini, analytics", "CLAUDE.md", "DECISION_LOG D-12", "G02, J04, R08", "Explicit owner approval", "Owner", "LOCKED"],
    ["D-26", "2026-10-03", "Client input never authoritative for measured state", "Gate 02", "Rules tightening; server chat storage", "Path-structural trust; authenticated Echo", "Forgeable browser state", "155cb69, 955c1cc", "A09, B02-B04, C02", "Open: client-asserted captureMethod/persistAttempt", "Owner", "IMPLEMENTED (branch); owner decision open"],
    ["D-21", "-", "Gate 01 stores raw analysis in scanAnalyses", "Raw preservation", "skinStates collection", "users/{uid}/scanAnalyses/{scanId}_{sha16}", "Immutable raw substrate", "6e9081d", "A06", "-", "Systems Architect", "IMPLEMENTED (not deployed); ID duplicated in HEAD"],
]

# ---------------------------------------------------------------- validation + output
ids = {r[0] for r in TASKS}
for r in TASKS:
    tid, status, deps, commit, ev, blocker, nxt = r[0], r[5], r[7], r[14], r[13], r[16], r[17]
    assert status in STATUSES, f"{tid}: bad status {status}"
    assert r[18] in MILESTONES, f"{tid}: bad milestone"
    assert nxt, f"{tid}: missing next action"
    if status == "DONE":
        assert commit or ev, f"{tid}: DONE without evidence"
    if status in ("BLOCKED", "OWNER ACTION", "RESEARCH"):
        assert blocker, f"{tid}: {status} without blocker"
    for d in filter(None, (x.strip() for x in deps.split(","))):
        assert d in ids, f"{tid}: unknown dependency {d}"
assert len(ids) == len(TASKS), "duplicate task IDs"

counts = {s: sum(r[5] == s for r in TASKS) for s in STATUSES}
total = len(TASKS)
in_scope = total - counts["LATER"]
pct = round(100 * counts["DONE"] / in_scope)
remaining = total - counts["DONE"]
core = [r for r in TASKS if r[5] != "LATER"]
milestone_done = {m: (sum(r[5] == "DONE" for r in core if r[18] == m), sum(r[18] == m for r in core)) for m in MILESTONES}
current_ms = next(m for m in MILESTONES if milestone_done[m][0] < milestone_done[m][1])

summary = [
    ("Overall execution percentage", f"{pct}% = DONE {counts['DONE']} / (total {total} - LATER {counts['LATER']}). Execution tracking only; NOT scientific accuracy or product readiness. DONE includes branch-only work not deployed."),
    ("Completed tasks (DONE)", counts["DONE"]), ("In-progress tasks", counts["IN PROGRESS"]),
    ("Blocked tasks", counts["BLOCKED"]), ("Owner-action tasks", counts["OWNER ACTION"]),
    ("Prepared/unblocked tasks", f"{counts['PREPARED'] + counts['UNBLOCKED']} (PREPARED {counts['PREPARED']}, UNBLOCKED {counts['UNBLOCKED']})"),
    ("Research-dependent tasks", counts["RESEARCH"]), ("Needs verification", counts["NEEDS VERIFICATION"]),
    ("Later tasks", counts["LATER"]), ("Remaining tasks (not DONE)", remaining),
    ("Current milestone", f"{current_ms} (first milestone with open tasks)"),
]
milestone_line = " -> ".join(f"{m} [{d}/{n}]" for m, (d, n) in milestone_done.items()) + f"  (excludes {counts['LATER']} LATER tasks)"


def md_cell(v):
    return str(v).replace("|", "\\|").replace("\n", " ")


def md_table(cols, rows):
    out = ["| " + " | ".join(cols) + " |", "|" + "---|" * len(cols)]
    out += ["| " + " | ".join(md_cell(c) for c in r) + " |" for r in rows]
    return "\n".join(out)


md = [
    "# PROJECT SKIN Execution Tracker",
    "",
    f"Single source of truth for project execution. Generated by `scripts/build_execution_tracker.py` on {UPDATED}; "
    "edit the script and regenerate, never edit this file or the XLSX by hand.",
    "",
    "Percentages measure execution tracking only. They are not scientific accuracy, measurement validity or product readiness.",
    "",
    "## 1. Overview",
    "",
    md_table(["Field", "Value"], OVERVIEW + summary),
    "",
    "### Milestones (DONE / core tasks)",
    "",
    milestone_line,
    "",
    "### Immediate next actions",
    "",
    *[f"{i}. {a}" for i, a in enumerate(IMMEDIATE_NEXT, 1)],
    "",
    "### Agent-executable now (read-only or design-only)",
    "",
    *[f"- {a}" for a in AGENT_NEXT],
    "",
    "### Critical blockers",
    "",
    *[f"- {b}" for b in CRITICAL_BLOCKERS],
    "",
    "### Discrepancies vs the seed baseline (verified 2026-10-03)",
    "",
    md_table(["Item", "Finding", "How verified"], DISCREPANCIES),
    "",
    "## 2. Master tracker",
    "",
    md_table(COLS, TASKS),
    "",
    "## 3. Research track",
    "",
    md_table(RESEARCH_COLS, R),
    "",
    "## 4. Claude operations",
    "",
    md_table(OPS_COLS, OPS),
    "",
    "### Reusable roles",
    "",
    md_table(["Role", "Scope"], ROLES),
    "",
    "## 5. Decisions & gates",
    "",
    "D-numbers are DECISION_LOG.md entries. `TD-` entries are tracker-local records of decisions found in gate reports but not in DECISION_LOG.",
    "",
    md_table(DEC_COLS, DECISIONS),
    "",
    "## 6. Operating rules",
    "",
    "1. This tracker is the project source of truth. 2. Never silently change architecture decisions. "
    "3. No DONE without evidence (commit, validated artifact, research report or explicit owner decision). "
    "4. Implementation completion is not production verification; research completion is not scientific validation. "
    "5. LATER stays LATER. 6. Agents never invent evidence. 7. Every BLOCKED task names its blocker; every task has a next action. "
    "8. Dependencies are explicit. 9. On completion, update this script and regenerate both files; record commit SHA and test/typecheck/lint/build/deploy evidence. "
    "10. Destructive production actions stop at OWNER ACTION. 11. No Terraform apply on the engine path without separate authorization. 12. No secrets here.",
    "",
]
MD_OUT.write_text("\n".join(md))

from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font  # noqa: E402

wb = Workbook()


def sheet(title, cols, rows, first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.append(cols)
    for r in rows:
        ws.append([str(c) if not isinstance(c, int) else c for c in r])
    for c in ws[1]:
        c.font = Font(bold=True)
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 28
        for c in col:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "B2"
    return ws


ov = OVERVIEW + summary + [("Milestones (DONE/total)", milestone_line)]
ov += [(f"Immediate next action {i}", a) for i, a in enumerate(IMMEDIATE_NEXT, 1)]
ov += [(f"Agent-executable {i}", a) for i, a in enumerate(AGENT_NEXT, 1)]
ov += [(f"Critical blocker {i}", b) for i, b in enumerate(CRITICAL_BLOCKERS, 1)]
ov += [(f"Discrepancy: {a}", f"{b} [{c}]") for a, b, c in DISCREPANCIES]
ws = sheet("Overview", ["Field", "Value"], ov, first=True)
ws.column_dimensions["B"].width = 120
sheet("Master Tracker", COLS, TASKS)
sheet("Research Track", RESEARCH_COLS, R)
ws = sheet("Claude Operations", OPS_COLS, OPS)
ws.append([])
ws.append(["Reusable role", "Scope"])
for role in ROLES:
    ws.append(list(role))
sheet("Decisions & Gates", DEC_COLS, DECISIONS)
wb.save(XLSX_OUT)

# read-back check
from openpyxl import load_workbook  # noqa: E402

wb2 = load_workbook(XLSX_OUT)
assert wb2.sheetnames == ["Overview", "Master Tracker", "Research Track", "Claude Operations", "Decisions & Gates"]
mt = list(wb2["Master Tracker"].iter_rows(min_row=2, values_only=True))
assert len(mt) == total and all(r[5] in STATUSES for r in mt)
md_lines = MD_OUT.read_text().splitlines()
assert all(any(line.startswith(f"| {tid} | ") for line in md_lines) for tid in ids), "MD missing task rows"
print(f"OK: {total} tasks, {pct}% execution ({counts}), current milestone {current_ms}")
